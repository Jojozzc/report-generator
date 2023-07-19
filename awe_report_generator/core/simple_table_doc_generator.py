from abc import ABCMeta, abstractmethod

from docx import Document
from typing import Dict, List

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt

from awe_report_generator.core.util import array_util
from awe_report_generator.core.util import cast_util
from .base import ReportGenerator, ExcelFiledProperty
from .style import DocCellStyle
from docx.oxml.ns import qn


class HeaderResource(metaclass=ABCMeta):
    """
    用于处理Word表格头数据
    """

    def set(self, doc: Document, data_list: List[Dict], global_data: dict):
        pass


class ColumnCellsParagraphAddRunResource():
    """
    用于处理word内表格内某一列数据
    """

    def __init__(self, table_index: int, table_data_start_row, column_view_index: int, paragraph_index=0,
                 style: DocCellStyle = DocCellStyle()):
        self.table_index = table_index
        self.table_data_start_row = table_data_start_row
        self.column_view_index = column_view_index
        self.style = style
        self.paragraph_index = paragraph_index

    def get_value(self, data: dict, global_data: dict, merged_data: dict):
        return None

    def set(self, cur_index: int, doc: Document, data: dict, global_data: dict,
            tables_row_index_zip: List[List[List[tuple]]], merged_data: dict):
        val = self.get_value(data, global_data, merged_data)
        if val is None:
            return
        table = doc.tables[self.table_index]
        row = self.table_data_start_row + cur_index

        index_zip = tables_row_index_zip[self.table_index]
        col = index_zip[row][self.column_view_index][0]
        cell = table.cell(row, col)
        para = cell.paragraphs[self.paragraph_index]
        if self.style is not None:
            para.paragraph_format.alignment = WD_TABLE_ALIGNMENT.CENTER

        if val is None:
            return

        font_cn = None

        if self.style is not None and self.style.font_cn is not None:
            font_cn = self.style.font_cn

        run = para.add_run(text=str(val))

        if font_cn is not None:
            run.font.name = font_cn
            run._element.rPr.rFonts.set(qn('w:eastAsia'), font_cn)

        if self.style is not None and self.style.font_size is not None:
            run.font.size = Pt(self.style.font_size)


class CellResource():
    """
    用于处理表格内某一个cell数据
    """

    def __init__(self, table_index: int, row: int, column_view_index: int, style: DocCellStyle = DocCellStyle(),
                 cast_value_to_str=cast_util.wrap_str):
        self.table_index = table_index
        self.row = row
        self.column_view_index = column_view_index
        self.style = style
        if cast_value_to_str is None:
            cast_value_to_str = cast_util.wrap_str
        self.cast_value_to_str = cast_value_to_str

    def get_value(self, data_list: List[dict], global_data: dict, merged_data: dict):
        return None

    def set(self, doc: Document, data_list: List[dict], global_data: dict,
            tables_row_index_zip: List[List[List[tuple]]], merged_data: dict):
        """

        :param tables_row_index_zip: M * x * y * 2 array
        :return:
        """
        val = self.get_value(data_list, global_data, merged_data)
        if val is None:
            return
        table = doc.tables[self.table_index]
        index_zip = tables_row_index_zip[self.table_index]

        col = index_zip[self.row][self.column_view_index][0]
        cell = table.cell(self.row, col)

        cell.text = self.cast_value_to_str(val)
        if self.style is not None and self.style.alignment is not None:
            cell.paragraphs[0].paragraph_format.alignment = self.style.alignment


class MappingColumnCellsParagraphAddRunResource(ColumnCellsParagraphAddRunResource):

    def __init__(self, table_index: int, table_data_start_row, column_view_index: int, mapping_data_key: str,
                 paragraph_index=0, style=DocCellStyle()):
        super().__init__(table_index, table_data_start_row, column_view_index, paragraph_index, style)
        self.mapping_data_key = mapping_data_key

    def get_value(self, data: dict, global_data: dict, merged_data: dict):
        if self.mapping_data_key in data:
            return data[self.mapping_data_key]
        return None


class CellParagraphAddRunResource(CellResource):
    def __init__(self, table_index: int, row: int, column_view_index: int, style: DocCellStyle = DocCellStyle(),
                 cast_value_to_str=cast_util.wrap_str, paragraph_index: int = 0, run_index: int = None):
        super().__init__(table_index, row, column_view_index, style, cast_value_to_str)
        self.table_index = table_index
        self.row = row
        self.column_view_index = column_view_index
        self.style = style
        self.paragraph_index = paragraph_index
        self.run_index = run_index

    def set(self, doc: Document, data_list: List[dict], global_data: dict,
            tables_row_index_zip: List[List[List[tuple]]], merged_data: dict):
        val = self.get_value(data_list, global_data, merged_data)
        if val is None:
            return
        table = doc.tables[self.table_index]
        index_zip = tables_row_index_zip[self.table_index]

        col = index_zip[self.row][self.column_view_index][0]
        cell = table.cell(self.row, col)
        para = cell.paragraphs[self.paragraph_index]

        val = self.cast_value_to_str(val)

        if val is None:
            return

        font_cn = None

        if self.style is not None and self.style.font_cn is not None:
            font_cn = self.style.font_cn
        if self.run_index is None:
            run = para.add_run(text=val)
        else:
            run = para.runs[self.run_index]
            run.text = val

        if font_cn is not None:
            run.font.name = font_cn
            run._element.rPr.rFonts.set(qn('w:eastAsia'), font_cn)

        if self.style is not None and self.style.font_size is not None:
            run.font.size = Pt(self.style.font_size)


class GlobalParamCellParagraphAddRunResource(CellParagraphAddRunResource):
    def __init__(self, table_index: int, row: int, column_view_index: int, style: DocCellStyle = DocCellStyle(),
                 paragraph_index: int = 0, cast_value_to_str=cast_util.wrap_str, mapping_key: str = None,
                 run_index: int = None):
        super().__init__(table_index, row, column_view_index, style, cast_value_to_str, paragraph_index, run_index)
        self.mapping_key = mapping_key

    def get_value(self, data_list: List[dict], global_data: dict, merged_data: dict):
        if global_data is None:
            return None
        return global_data.get(self.mapping_key, None)


class DataListCellParagraphAddRunResource(CellParagraphAddRunResource):
    def __init__(self, table_index: int, row: int, column_view_index: int, style: DocCellStyle = DocCellStyle(),
                 paragraph_index: int = 0, cast_value_to_str=cast_util.wrap_str, mapping_key: str = None,
                 run_index: int = None, data_list_sort_func=None, data_list_sort_reverse=False):
        super().__init__(table_index, row, column_view_index, style, cast_value_to_str, paragraph_index, run_index)
        self.mapping_key = mapping_key
        self.data_list_sort_func = data_list_sort_func
        self.data_list_sort_reverse = data_list_sort_reverse

    def get_value(self, data_list: List[dict], global_data: dict, merged_data: dict):
        if self.data_list_sort_func is not None:
            data_list = sorted(data_list, key=lambda data: self.data_list_sort_func(data[self.mapping_key]), reverse=self.data_list_sort_reverse)
        return array_util.get_one_value(data_list, self.mapping_key)


class MergedDataCellParagraphAddRunResource(CellParagraphAddRunResource):
    def __init__(self, table_index: int, row: int, column_view_index: int, style: DocCellStyle = DocCellStyle(),
                 paragraph_index: int = 0, cast_value_to_str=cast_util.wrap_str, mapping_key: str = None,
                 run_index: int = None):
        super().__init__(table_index, row, column_view_index, style, cast_value_to_str, paragraph_index, run_index)
        self.mapping_key = mapping_key

    def get_value(self, data_list: List[dict], global_data: dict, merged_data: dict):
        if merged_data is None:
            return None
        return merged_data.get(self.mapping_key, None)


class SimpleCalculationColumnCellsParagraphAddRunResource(ColumnCellsParagraphAddRunResource):
    def __init__(self, table_index: int, table_data_start_row, column_view_index: int, paragraph_index: int = 0,
                 style: DocCellStyle = DocCellStyle(), mapping_data_key_1: str = None, mapping_data_key_2: str = None,
                 operation: str = None):
        """

        :param mapping_data_key_1: must be key of number
        :param mapping_data_key_2: must be key of number
        :param operation: +/-
        """
        super().__init__(table_index, table_data_start_row, column_view_index, paragraph_index, style)
        self.mapping_data_key_1 = mapping_data_key_1
        self.mapping_data_key_2 = mapping_data_key_2
        self.operation = operation

    def get_value(self, data: dict, global_data: dict, merged_data: dict):
        if global_data is None:
            return None

        if self.mapping_data_key_1 not in data:
            return None

        if self.mapping_data_key_2 not in data:
            return None

        if self.operation == '+':
            return data[self.mapping_data_key_1] + data[self.mapping_data_key_2]
        elif self.operation == '-':
            return data[self.mapping_data_key_1] - data[self.mapping_data_key_2]
        else:
            raise ValueError(f'Unknown operation:{self.operation}')


class SimpleTableDocGenerator(ReportGenerator):

    def __init__(self, template_path: str, filed_mapping: Dict[str, ExcelFiledProperty], divide_key: str,
                 header_resource: HeaderResource, column_cell_resource_list: List[ColumnCellsParagraphAddRunResource],
                 cell_resource_list: List[CellResource], doc_global_data_param_config_list=None,
                 merge_fun_dict: dict = None):
        super().__init__(template_path, filed_mapping, divide_key, doc_global_data_param_config_list, merge_fun_dict)
        self.header_resource = header_resource
        self.column_cell_resource_list = column_cell_resource_list
        self.cell_resource_list = cell_resource_list
        template_doc = Document(template_path)
        self.table_row_index_zips = self._build_template_doc_table_index_zips(template_doc)

    def _process(self, data_list: list, template_doc: Document, global_param: dict, merged_data: dict) -> Document:
        if self.header_resource is not None:
            self.header_resource.set(doc=template_doc, data_list=data_list, global_data=global_param)
        if self.column_cell_resource_list is not None:
            for i in range(len(data_list)):
                data = data_list[i]
                for col_res in self.column_cell_resource_list:
                    col_res.set(i, template_doc, data, global_param, self.table_row_index_zips, merged_data)
        if self.cell_resource_list is not None:
            for cell_res in self.cell_resource_list:
                cell_res.set(template_doc, data_list, global_param, self.table_row_index_zips, merged_data)

        return template_doc

    def _build_template_doc_table_index_zips(self, doc: Document):
        """
        :return M * x * y * 2 array:
            M = len of table.tables
            x = count of rows/columns in one table
            y = count of block
            2 = (start, end) end is excluded, start < end
            [M=2 tables
                [table0:x=3 rows/columns
                    [(0, 2),(2, 6),(6, 7)],
                    [(0, 3),(3, 7)],
                    [(0, 2),(2, 5), (5, 7)],
                ],
                [table1:x=2 rows/columns
                    [(0, 3),(3, 8)],
                    [(0, 2),(2, 8)],
                ]
            ]
        """
        tables = doc.tables
        index_zips = []
        if tables is None:
            return index_zips

        for table in tables:
            zips = []
            index_zips.append(zips)

            row_or_column_cnt = len(table.rows)

            for i in range(row_or_column_cnt):
                cells = table.row_cells(i)
                zips.append(array_util.zip_arr(cells))

        return index_zips
