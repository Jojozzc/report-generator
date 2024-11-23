from abc import ABC
from typing import List, Dict

from _decimal import Decimal
from docx import Document

from awe_report_generator.core.simple_table_doc_generator import HeaderResource, CellResource, \
    CellParagraphAddRunResource
from awe_report_generator.core.util import array_util, cast_util


class RTHeaderResource(HeaderResource, ABC):

    def set(self, doc: Document, data_list: List[Dict], global_data: dict):
        project_name = global_data.get('projectName', '')
        if project_name is not None:
            doc.paragraphs[1].add_run(project_name)
        unit_name = array_util.get_one_value(data_list, 'unitName')
        if unit_name is not None:
            doc.paragraphs[2].add_run(unit_name)



class RayAdditionalHeaderResource(HeaderResource, ABC):

    def set(self, doc: Document, data_list: List[Dict], global_data: dict):
        project_name = global_data.get('projectName', '')
        customorCompany = global_data.get('customorCompany', '')
        doc.paragraphs[2].add_run(f'工程名称：{project_name}                         委托单位：{customorCompany}')




class RecordHeaderResource(HeaderResource, ABC):

    def set(self, doc: Document, data_list: List[Dict], global_data: dict):
        project_name = global_data.get('projectName', '')
        if project_name is not None:
            doc.paragraphs[1].add_run(project_name)
        unit_name = global_data.get('customerCompany', '')
        if unit_name is not None:
            doc.paragraphs[2].add_run(unit_name)


class RTSummaryCellResource(CellResource):
    SUMMARY_FORMAT_ONE = '说明：共检测{data_size}道，合格{ok_data_size}道，不合格{bad_data_size}道，其中返修{bad_check_count}张，共计{check_count}张。'
    SUMMARY_FORMAT_TWO = '其中γ射线{ray}张。'

    def get_value(self, data_list: List[dict], global_data: dict, merged_data: dict):
        if data_list is None:
            return None
        data_size = len(data_list)
        ok_data_size = 0

        check_count = 0
        ok_check_count = 0
        ray = 0

        for data in data_list:
            check_c = data.get('checkCount', 0)
            check_count += check_c

            ok_check_c = data.get('okCount', 0)
            ok_check_count += ok_check_c

            if check_c == ok_check_c:
                ok_data_size += 1

            ray_str = data.get('ray', '')
            if ray_str is not None and ray_str == 'γ':
                ray += check_c

        val = self.SUMMARY_FORMAT_ONE.format(data_size=data_size, ok_data_size=ok_data_size,
                                             bad_data_size=data_size - ok_data_size,
                                             bad_check_count=check_count - ok_check_count, check_count=check_count)

        if ray > 0:
            val = val + self.SUMMARY_FORMAT_TWO.format(ray=ray)

        return val


class RTRecordSummaryCellParagraphAddRunResource(CellParagraphAddRunResource):
    SUMMARY_FORMAT = '检测焊口（检件）数  {data_size} 道     底片数 {check_count} 张'

    def get_value(self, data_list: List[dict], global_data: dict, merged_data: dict):
        if data_list is None:
            return None
        data_size = len(data_list)

        check_count = 0


        for data in data_list:
            check_c = data.get('checkCount', 0)
            check_count += check_c

        val = self.SUMMARY_FORMAT.format(data_size=data_size, check_count=check_count)

        return val


class RTSummaryCellResource2(CellResource):
    SUMMARY_FORMAT_ONE = '说明：共检测{data_size}道，合格{ok_data_size}道，不合格{bad_data_size}道'

    def get_value(self, data_list: List[dict], global_data: dict, merged_data: dict):
        if data_list is None:
            return None
        data_size = len(data_list)
        ok_data_size = 0

        check_count = 0
        ok_check_count = 0

        meter_sum = Decimal('0')
        line_sum = Decimal('0')

        for data in data_list:
            check_c = data.get('checkCount', 0)
            check_count += check_c

            ok_check_c = data.get('okCount', 0)
            ok_check_count += ok_check_c

            if check_c == ok_check_c:
                ok_data_size += 1

            detection_count: str = data.get('detectionCount', None)
            if detection_count is None:
                pass
            else:
                detection_count_num = cast_util.check_ret_float_str(detection_count[:-1])
                if detection_count_num is not None:
                    if detection_count.endswith('m'):
                        meter_sum = meter_sum + Decimal(detection_count_num)
                    elif detection_count.endswith('道'):
                        line_sum += line_sum + Decimal(detection_count_num)

        agg_val: str
        if meter_sum <= 0 and line_sum <= 0:
            agg_val = f'共计{data_size}道'
        elif meter_sum <= 0:
            agg_val = f'共计{cast_util.wrap_int_str(line_sum)}道'
        elif line_sum <= 0:
            agg_val = f'共计{meter_sum}米'
        else:
            agg_val = f'共计{cast_util.wrap_int_str(line_sum)}道，{meter_sum}米'
        val = self.SUMMARY_FORMAT_ONE.format(data_size=data_size, ok_data_size=ok_data_size,
                                             bad_data_size=data_size - ok_data_size) + '，' + agg_val

        return val
