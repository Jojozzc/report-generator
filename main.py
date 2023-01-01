import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
import pandas
import os

def write_cell(cell, text:str):
    cell.text = text
    cell.paragraphs[0].paragraph_format.alignment = WD_TABLE_ALIGNMENT.CENTER


if __name__ == '__main__':
    # file_path = './testData/EPK-BFJC-MM-3115-AZ(02)-RT-039.docx'
    input_file_path = './testData/气化数据库2.xlsx'
    template_path = './template/TEMPLATE.docx'


    raw_data = pandas.read_excel(input_file_path, 'Sheet1')

    data_map = {}

    for i in range(raw_data.shape[0]):
        index = raw_data.iloc[i, 2]
        if index not in data_map:
            data_map[index] = []
        data_map[index].append(raw_data.iloc[i])



    for key in data_map.keys():
        values = data_map[key]
        order_id = values[0][2]
        quality_level = values[0][8]

        doc = docx.Document(template_path)

        target_path = os.path.join('./testData', f'{order_id}.docx')

        doc.paragraphs[1].text = doc.paragraphs[1].text + '宁夏宝丰能源集团股份有限公司50万吨/年煤制烯烃项目配套甲醇工程'
        doc.paragraphs[2].text = doc.paragraphs[2].text + '气化1'
        table = doc.tables[0]

        write_cell(table.cell(0, 2), '南京英派克检测有限责任公司')
        write_cell(table.cell(0, 11), order_id)

        write_cell(table.cell(1, 2), '中化六建')
        write_cell(table.cell(1, 11), '2022.12.02')

        write_cell(table.cell(2, 2), 'RT')
        write_cell(table.cell(2, 8), 'NB/T47013.2-2015')

        write_cell(table.cell(2, 14), format(quality_level, '.0%'))


        kind_count = 0
        unqualified_kind_cnt = 0

        sample_cnt = 0
        unqualified_sample_cnt = 0


        # 数据从第5行开始写入
        for i in range(len(values)):
            row = values[i]
            # 检件编号
            table.cell(i + 5, 1).text = row[3]
            table.cell(i + 5, 1).paragraphs[0].paragraph_format.alignment = WD_TABLE_ALIGNMENT.CENTER

            # 焊口编号
            table.cell(i + 5, 3).text = str(row[4])
            table.cell(i + 5, 3).paragraphs[0].paragraph_format.alignment = WD_TABLE_ALIGNMENT.CENTER


            # 材质
            table.cell(i + 5, 5).text = str(row[7])
            table.cell(i + 5, 5).paragraphs[0].paragraph_format.alignment = WD_TABLE_ALIGNMENT.CENTER


            # 规格
            table.cell(i + 5, 7).text = str(row[6])
            table.cell(i + 5, 7).paragraphs[0].paragraph_format.alignment = WD_TABLE_ALIGNMENT.CENTER

            # 底片规格、数量
            sample_cnt = sample_cnt + row[11]
            table.cell(i + 5, 9).text = f'{row[10]}/{row[11]}张'
            table.cell(i + 5, 9).paragraphs[0].paragraph_format.alignment = WD_TABLE_ALIGNMENT.CENTER

            # 检测结果(合格)
            table.cell(i + 5, 13).text = f'{row[12]}张'
            table.cell(i + 5, 13).paragraphs[0].paragraph_format.alignment = WD_TABLE_ALIGNMENT.CENTER

            # 检测结果(不合格)
            unqualified_cnt = row[11] - row[12]
            if unqualified_cnt > 0:
                unqualified_kind_cnt = unqualified_kind_cnt + 1
                table.cell(i + 5, 15).text = f'{unqualified_cnt}张'
            else:
                table.cell(i + 5, 15).text = '/'

            unqualified_sample_cnt = unqualified_sample_cnt + unqualified_cnt
            table.cell(i + 5, 15).paragraphs[0].paragraph_format.alignment = WD_TABLE_ALIGNMENT.CENTER
            kind_count = kind_count + 1

        if kind_count < 18:
            table.cell(kind_count + 5, 1).text = '以下空白'
            cel = table.cell(kind_count + 5, 1)
            table.cell(kind_count + 5, 1).paragraphs[0].paragraph_format.alignment = WD_TABLE_ALIGNMENT.CENTER

        table.cell(18 + 5, 1).text = f'说明：共检测焊口{kind_count}道口，总计{sample_cnt}张底片。其中不合格焊{unqualified_kind_cnt}道，不合格底片{unqualified_sample_cnt}张。'

        doc.save(target_path)