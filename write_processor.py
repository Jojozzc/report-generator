import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
import pandas
import os
from RawData import RawData


def write_cell(cell, text: str):
    cell.text = text
    cell.paragraphs[0].paragraph_format.alignment = WD_TABLE_ALIGNMENT.CENTER


def process(template_path: str, input_file_path: str, title1: str, title2: str, company_name: str, customer: str,
            method: str, standard: str, target_dir: str):
    base_start = 5
    raw_data_map = read_raw_data(input_file_path)

    for key in raw_data_map.keys():
        try:
            values = raw_data_map[key]
            order_id = key
            quality_level = values[0].quality_level
            complete_date = values[0].complete_date

            doc = docx.Document(template_path)

            target_path = os.path.join(target_dir, f'{order_id}.docx')

            doc.paragraphs[1].text = doc.paragraphs[1].text + title1
            doc.paragraphs[2].text = doc.paragraphs[2].text + title2
            table = doc.tables[0]

            write_cell(table.cell(0, 2), company_name)
            write_cell(table.cell(0, 11), order_id)

            write_cell(table.cell(1, 2), customer)
            write_cell(table.cell(1, 11), complete_date)

            write_cell(table.cell(2, 2), method)
            write_cell(table.cell(2, 8), standard)

            write_cell(table.cell(2, 14), format(quality_level, '.0%'))

            kind_count = 0
            unqualified_kind_cnt = 0

            sample_cnt = 0
            unqualified_sample_cnt = 0

            # 数据从第5行开始写入
            for i in range(len(values)):
                raw_data: RawData = values[i]

                # 检件编号
                write_cell(table.cell(i + base_start, 1), raw_data.sample_no)

                # 焊口编号
                write_cell(table.cell(i + base_start, 3), raw_data.kind_no)

                # 材质
                write_cell(table.cell(i + base_start, 5), raw_data.material)

                # 规格
                write_cell(table.cell(i + base_start, 7), raw_data.specification)

                sample_specification = ''
                if raw_data.sample_specification is not None:
                    sample_specification = raw_data.sample_specification
                # 底片规格、数量
                if raw_data.sample_cnt is not None:
                    sample_cnt = sample_cnt + raw_data.sample_cnt
                    write_cell(table.cell(i + base_start, 9), f'{sample_specification}/{raw_data.sample_cnt}张')
                else:
                    print(f'数据有误(张数)请检查：委托单号={order_id},检件编号={raw_data.sample_no}')
                    write_cell(table.cell(i + base_start, 9), f'{sample_specification}')

                # 检测结果(合格)
                if raw_data.qualified_sample_cnt is not None:
                    write_cell(table.cell(i + base_start, 13), f'{raw_data.qualified_sample_cnt}张')
                else:
                    write_cell(table.cell(i + base_start, 13), '/')

                # 检测结果(不合格)
                if raw_data.sample_cnt is not None and raw_data.qualified_sample_cnt is not None:
                    unqualified_cnt = raw_data.sample_cnt - raw_data.qualified_sample_cnt
                    if unqualified_cnt > 0:
                        unqualified_sample_cnt = unqualified_sample_cnt + unqualified_cnt
                        unqualified_kind_cnt = unqualified_kind_cnt + 1
                        write_cell(table.cell(i + base_start, 15), f'{unqualified_cnt}张')
                    else:
                        write_cell(table.cell(i + base_start, 15), '/')
                else:
                    print(f'数据有误(张数或合格数量)请检查：委托单号={order_id},检件编号={raw_data.sample_no}')

                kind_count = kind_count + 1

            if kind_count < 22:
                write_cell(table.cell(kind_count + base_start, 1), '以下空白')

            table.cell(22 + base_start,
                       1).text = f'说明：共检测焊口{kind_count}道口，总计{sample_cnt}张底片。其中不合格焊{unqualified_kind_cnt}道，不合格底片{unqualified_sample_cnt}张。'

            doc.save(target_path)
        except Exception as e:
            print(f'数据有误,请检查：委托单号={key}')
            print(e)


def read_raw_data(input_file_path: str):
    raw_datas = pandas.read_excel(input_file_path, 'Sheet1')
    data_map = {}

    for i in range(raw_datas.shape[0]):
        row = raw_datas.iloc[i]
        order_id = row[2]
        if order_id not in data_map:
            data_map[order_id] = []

        sample_cnt = None
        qualified_sample_cnt = None

        if str(row[11]).isdigit():
            sample_cnt = int(row[11])

        if str(row[12]).isdigit():
            qualified_sample_cnt = int(row[12])

        raw_data = RawData(wrapper_str(row[1]), wrapper_str(row[3]), wrapper_str(row[4]), wrapper_str(row[7]), wrapper_str(row[6]), wrapper_str(row[10]), sample_cnt,
                           qualified_sample_cnt, row[8])
        data_map[order_id].append(raw_data)

    return data_map


def wrapper_str(obj):
    if str(obj) == 'nan' or str(obj) == None:
        return None
    return str(obj)
